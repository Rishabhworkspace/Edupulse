"""EduPulse Project Presentation Generator"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
BG_DARK = RGBColor(0x0F, 0x17, 0x2A)
BG_CARD = RGBColor(0x1A, 0x24, 0x3B)
CORAL = RGBColor(0xFF, 0x6B, 0x6B)
TEAL = RGBColor(0x4E, 0xCE, 0xC6)
GOLD = RGBColor(0xFF, 0xD9, 0x3D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xA0, 0xAE, 0xC1)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
GREEN = RGBColor(0x6B, 0xCB, 0x77)
BLUE = RGBColor(0x45, 0x6E, 0xFF)

def dark_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

def add_text(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_shape(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_accent_line(slide, left, top, width, color=CORAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.06))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def card(slide, left, top, w, h, title, bullets, accent=CORAL, icon=""):
    add_shape(slide, left, top, w, h, BG_CARD)
    add_accent_line(slide, left, top, w, accent)
    label = f"{icon}  {title}" if icon else title
    add_text(slide, left+0.3, top+0.2, w-0.6, 0.5, label, size=16, color=accent, bold=True)
    for i, b in enumerate(bullets):
        add_text(slide, left+0.3, top+0.75+i*0.38, w-0.6, 0.4, f"•  {b}", size=12, color=GRAY)

# ============================================================
# SLIDE 1 — Title
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 3.5, 2.3, 6.3, CORAL)
add_text(s, 1, 2.5, 11.3, 1.2, "EduPulse", size=64, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, 1, 3.7, 11.3, 0.6, "A Modern Learning Management System", size=24, color=TEAL, align=PP_ALIGN.CENTER)
add_text(s, 1, 4.5, 11.3, 0.5, "Full-Stack MERN Application  •  Built by Rishabh", size=16, color=GRAY, align=PP_ALIGN.CENTER)
add_accent_line(s, 5, 5.3, 3.3, TEAL)

# ============================================================
# SLIDE 2 — Agenda
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 2, CORAL)
add_text(s, 0.8, 0.9, 5, 0.7, "Agenda", size=36, color=WHITE, bold=True)
items = [
    ("01", "Project Overview", CORAL),
    ("02", "Tech Stack & Architecture", TEAL),
    ("03", "Student Features", PURPLE),
    ("04", "Instructor & Admin Features", GOLD),
    ("05", "Database Design", GREEN),
    ("06", "API Layer & Security", BLUE),
    ("07", "CI/CD & Deployment", CORAL),
    ("08", "Future Roadmap", TEAL),
]
for i, (num, title, color) in enumerate(items):
    y = 2.0 + i * 0.6
    add_shape(s, 1.5, y, 10, 0.48, BG_CARD)
    add_text(s, 1.7, y+0.05, 0.6, 0.4, num, size=16, color=color, bold=True)
    add_text(s, 2.4, y+0.05, 8, 0.4, title, size=16, color=WHITE)

# ============================================================
# SLIDE 3 — Project Overview
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 2.5, CORAL)
add_text(s, 0.8, 0.9, 6, 0.7, "Project Overview", size=36, color=WHITE, bold=True)
add_text(s, 0.8, 1.8, 11.5, 0.8,
    "EduPulse is a comprehensive LMS enabling educators to publish courses and students to learn interactively. "
    "It features role-based access, real-time progress tracking, Stripe payments, and a modern React UI.",
    size=14, color=GRAY)
card(s, 0.8, 3.0, 3.6, 2.8, "Students", [
    "Browse & enroll in courses", "Video + article lessons",
    "Track progress & XP", "Take notes per lesson",
    "Earn certificates"], CORAL, "🎓")
card(s, 4.8, 3.0, 3.6, 2.8, "Instructors", [
    "Course Studio builder", "Curriculum management",
    "Revenue analytics", "Q&A moderation",
    "Student engagement"], TEAL, "📊")
card(s, 8.8, 3.0, 3.6, 2.8, "Admins", [
    "User & role management", "Course moderation",
    "Coupon system", "Order tracking",
    "Platform analytics"], GOLD, "🔧")

# ============================================================
# SLIDE 4 — Tech Stack
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 2, TEAL)
add_text(s, 0.8, 0.9, 6, 0.7, "Tech Stack", size=36, color=WHITE, bold=True)
card(s, 0.8, 2.0, 3.7, 2.5, "Frontend", [
    "React 18 + Vite", "Redux Toolkit (state)",
    "Framer Motion (animation)", "Lucide React (icons)",
    "React Router v6"], CORAL, "⚛️")
card(s, 4.8, 2.0, 3.7, 2.5, "Backend", [
    "Node.js + Express.js", "MongoDB + Mongoose ODM",
    "JWT dual-token auth", "Passport.js (Google OAuth)",
    "RESTful API v1"], TEAL, "🖥️")
card(s, 8.8, 2.0, 3.7, 2.5, "Infrastructure", [
    "Cloudinary (media CDN)", "Stripe (payments)",
    "GitHub Actions CI/CD", "Vercel (deployment)",
    "MongoDB Atlas (cloud DB)"], PURPLE, "☁️")
card(s, 0.8, 4.8, 5.8, 1.6, "Security Stack", [
    "Helmet HTTP headers  •  Rate limiting  •  Mongo sanitize  •  XSS clean"], GOLD, "🔒")
card(s, 7.0, 4.8, 5.5, 1.6, "Dev Tooling", [
    "npm workspaces (monorepo)  •  Concurrently  •  Vitest  •  Jest"], GREEN, "🛠️")

# ============================================================
# SLIDE 5 — Architecture
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 2.5, TEAL)
add_text(s, 0.8, 0.9, 8, 0.7, "System Architecture", size=36, color=WHITE, bold=True)

# Client box
add_shape(s, 0.8, 2.2, 3.5, 4.0, BG_CARD)
add_accent_line(s, 0.8, 2.2, 3.5, CORAL)
add_text(s, 1.0, 2.4, 3, 0.4, "🌐  Client (Port 5173)", size=14, color=CORAL, bold=True)
for i, t in enumerate(["React 18 + Vite", "Redux Toolkit Store", "Lazy-loaded Routes", "Axios Interceptors", "ErrorBoundary", "Protected Routes"]):
    add_text(s, 1.1, 2.9+i*0.38, 3, 0.35, f"▸ {t}", size=11, color=GRAY)

# Arrow
add_shape(s, 4.5, 3.8, 1.2, 0.08, TEAL)
add_text(s, 4.5, 3.4, 1.2, 0.4, "/api proxy", size=10, color=TEAL, align=PP_ALIGN.CENTER)

# Server box
add_shape(s, 5.9, 2.2, 3.5, 4.0, BG_CARD)
add_accent_line(s, 5.9, 2.2, 3.5, TEAL)
add_text(s, 6.1, 2.4, 3, 0.4, "⚙️  Server (Port 5000)", size=14, color=TEAL, bold=True)
for i, t in enumerate(["Express.js + Router", "JWT Auth Middleware", "Rate Limiter", "Mongoose ODM", "Validators", "Error Handler"]):
    add_text(s, 6.2, 2.9+i*0.38, 3, 0.35, f"▸ {t}", size=11, color=GRAY)

# Arrow 2
add_shape(s, 9.6, 3.8, 1.2, 0.08, GOLD)
add_text(s, 9.6, 3.4, 1.2, 0.4, "Mongoose", size=10, color=GOLD, align=PP_ALIGN.CENTER)

# DB box
add_shape(s, 11.0, 2.8, 1.8, 2.0, BG_CARD)
add_accent_line(s, 11.0, 2.8, 1.8, GOLD)
add_text(s, 11.1, 3.0, 1.6, 0.4, "🗄️ MongoDB", size=13, color=GOLD, bold=True)
add_text(s, 11.1, 3.5, 1.6, 0.35, "Atlas Cloud", size=11, color=GRAY)
add_text(s, 11.1, 3.9, 1.6, 0.35, "10 Collections", size=11, color=GRAY)

# ============================================================
# SLIDE 6 — Student Features
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 3, PURPLE)
add_text(s, 0.8, 0.9, 8, 0.7, "Student Experience", size=36, color=WHITE, bold=True)
card(s, 0.8, 2.0, 3.7, 2.2, "Course Discovery", [
    "Searchable catalog", "Category filtering",
    "Rating & reviews", "Free & paid courses"], CORAL, "🔍")
card(s, 4.8, 2.0, 3.7, 2.2, "Learning Player", [
    "HD video playback", "Article viewer",
    "Quiz engine", "Personal notes"], TEAL, "▶️")
card(s, 8.8, 2.0, 3.7, 2.2, "Progress Tracking", [
    "Lesson completion %", "XP & streak system",
    "Schedule planner", "Certificate generation"], PURPLE, "📈")
card(s, 0.8, 4.5, 5.8, 2.0, "Community & Engagement", [
    "Course Q&A discussions with instructor replies",
    "Community forum for peer interaction",
    "Assignment submission & tracking",
    "Saved courses bookmark list"], GREEN, "💬")
card(s, 7.0, 4.5, 5.5, 2.0, "Dashboard Modules", [
    "StudentDashboard — overview with stats & schedule",
    "MyCoursesPage — enrolled course grid",
    "AchievementsPage — XP, badges, streaks",
    "AssignmentsPage — pending & completed work"], BLUE, "📋")

# ============================================================
# SLIDE 7 — Instructor & Admin
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 4, GOLD)
add_text(s, 0.8, 0.9, 10, 0.7, "Instructor & Admin Panels", size=36, color=WHITE, bold=True)
card(s, 0.8, 2.0, 5.8, 2.5, "Instructor Dashboard", [
    "Course Studio — create & edit full curricula",
    "Add video, article, and quiz lessons",
    "Revenue & enrollment analytics charts",
    "Q&A Dashboard — respond to student questions",
    "Manage course pricing & thumbnails"], TEAL, "🎬")
card(s, 7.0, 2.0, 5.5, 2.5, "Admin Controls", [
    "AdminUsersPage — manage all user accounts & roles",
    "AdminCoursesPage — moderate the entire catalog",
    "AdminOrdersPage — track payments & refunds",
    "AdminCouponsPage — create flat/% discount codes",
    "Role-based access: Student → Instructor → Admin"], GOLD, "🏢")
card(s, 0.8, 4.8, 11.7, 1.6, "Role-Based Access Control (RBAC)", [
    "JWT access + refresh token pair  •  Google OAuth via Passport.js  •  Middleware guards per route",
    "RoleRoute component on frontend  •  ProtectedRoute + GuestRoute wrappers  •  Auto token refresh"], PURPLE, "🔐")

# ============================================================
# SLIDE 8 — Database Design
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 3, GREEN)
add_text(s, 0.8, 0.9, 8, 0.7, "Database Design", size=36, color=WHITE, bold=True)
add_text(s, 0.8, 1.7, 11, 0.4, "MongoDB Atlas  •  10 Mongoose Collections  •  Optimized with .lean() queries", size=14, color=GRAY)
models = [
    ("User", "name, email, password, role, avatar, verified", CORAL),
    ("Course", "title, slug, price, instructor, thumbnail, rating", TEAL),
    ("Lesson", "title, type (video/article/quiz), content, order", PURPLE),
    ("Enrollment", "user, course, progress, completedLessons", GREEN),
    ("Order", "user, course, amount, status, paymentId", GOLD),
    ("Discussion", "course, user, question, replies, votes", BLUE),
    ("Review", "user, course, rating, comment", CORAL),
    ("Coupon", "code, type (flat/percent), discount, expiry", TEAL),
    ("Category", "name, slug, courseCount", PURPLE),
    ("Contact", "name, email, subject, message", GREEN),
]
for i, (name, fields, color) in enumerate(models):
    row = i // 2
    col = i % 2
    x = 0.8 + col * 6.2
    y = 2.3 + row * 0.9
    add_shape(s, x, y, 5.8, 0.75, BG_CARD)
    add_text(s, x+0.2, y+0.08, 1.5, 0.35, name, size=13, color=color, bold=True)
    add_text(s, x+1.8, y+0.08, 4, 0.6, fields, size=10, color=GRAY)

# ============================================================
# SLIDE 9 — API & Security
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 3, BLUE)
add_text(s, 0.8, 0.9, 8, 0.7, "API Layer & Security", size=36, color=WHITE, bold=True)
card(s, 0.8, 2.0, 5.8, 4.2, "REST API v1 Endpoints", [
    "POST /api/v1/auth/register — signup + email verify",
    "POST /api/v1/auth/login — JWT token pair",
    "GET  /api/v1/courses — catalog with filters",
    "GET  /api/v1/courses/:slug — full course detail",
    "POST /api/v1/payment/create — Stripe checkout",
    "POST /api/v1/payment/verify — confirm payment",
    "GET  /api/v1/users/me — current user profile",
    "CRUD /api/v1/discussions — course Q&A",
    "CRUD /api/v1/coupons — admin coupon mgmt"], BLUE, "🌐")
card(s, 7.0, 2.0, 5.5, 2.0, "Security Measures", [
    "Helmet — secure HTTP headers",
    "express-rate-limit — DDoS protection",
    "express-mongo-sanitize — injection prevention",
    "xss-clean — script injection blocking"], CORAL, "🛡️")
card(s, 7.0, 4.3, 5.5, 1.9, "Auth Architecture", [
    "Access token (15m) + Refresh token (7d)",
    "HttpOnly cookies for refresh tokens",
    "Axios interceptor auto-refresh",
    "Google OAuth 2.0 via Passport.js"], GREEN, "🔑")

# ============================================================
# SLIDE 10 — CI/CD
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 3, CORAL)
add_text(s, 0.8, 0.9, 8, 0.7, "CI/CD & Deployment", size=36, color=WHITE, bold=True)
steps = [
    ("1", "Push", "Developer pushes to main branch", CORAL),
    ("2", "Lint", "ESLint checks code quality", TEAL),
    ("3", "Test", "Jest (server) + Vitest (client)", PURPLE),
    ("4", "Build", "Vite production bundle", GOLD),
    ("5", "Deploy", "Auto-deploy to Vercel", GREEN),
]
for i, (num, label, desc, color) in enumerate(steps):
    x = 0.8 + i * 2.4
    add_shape(s, x, 2.5, 2.1, 1.8, BG_CARD)
    add_accent_line(s, x, 2.5, 2.1, color)
    add_text(s, x+0.2, 2.7, 1.7, 0.4, f"{num}. {label}", size=16, color=color, bold=True)
    add_text(s, x+0.2, 3.2, 1.7, 0.8, desc, size=11, color=GRAY)
    if i < 4:
        add_shape(s, x+2.15, 3.3, 0.2, 0.06, GRAY)

card(s, 0.8, 4.8, 5.8, 1.6, "Monorepo Structure", [
    "npm workspaces: client/ + server/",
    "Root scripts: dev, build, test",
    "Shared config via package.json"], TEAL, "📦")
card(s, 7.0, 4.8, 5.5, 1.6, "Environment Config", [
    "MONGO_URI, JWT secrets, Stripe keys",
    "VITE_API_URL for client proxy",
    ".env.example for onboarding"], GOLD, "⚙️")

# ============================================================
# SLIDE 11 — Future Roadmap
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 0.8, 0.7, 3, TEAL)
add_text(s, 0.8, 0.9, 8, 0.7, "Future Roadmap", size=36, color=WHITE, bold=True)
roadmap = [
    ("Phase 1", "Live Sessions", "WebRTC video classes + real-time chat", CORAL),
    ("Phase 2", "AI Tutor", "AI-powered learning assistant & quiz generation", TEAL),
    ("Phase 3", "Mobile App", "React Native companion app with offline support", PURPLE),
    ("Phase 4", "Marketplace", "Instructor payouts, affiliate system, analytics", GOLD),
]
for i, (phase, title, desc, color) in enumerate(roadmap):
    y = 2.2 + i * 1.2
    add_shape(s, 1.5, y, 10.3, 0.95, BG_CARD)
    add_accent_line(s, 1.5, y, 0.08, color)
    add_text(s, 1.8, y+0.1, 1.5, 0.35, phase, size=12, color=color, bold=True)
    add_text(s, 3.3, y+0.1, 3, 0.35, title, size=16, color=WHITE, bold=True)
    add_text(s, 3.3, y+0.5, 8, 0.35, desc, size=12, color=GRAY)

# ============================================================
# SLIDE 12 — Thank You
# ============================================================
s = prs.slides.add_slide(prs.slide_layouts[6])
dark_bg(s)
add_accent_line(s, 4, 2.3, 5.3, CORAL)
add_text(s, 1, 2.5, 11.3, 1, "Thank You!", size=56, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, 1, 3.6, 11.3, 0.6, "Questions & Discussion", size=24, color=TEAL, align=PP_ALIGN.CENTER)
add_accent_line(s, 5, 4.5, 3.3, TEAL)
add_text(s, 1, 5.0, 11.3, 0.5, "Built with ❤️ by Rishabh  •  github.com/Rishabhworkspace/Edupulse", size=14, color=GRAY, align=PP_ALIGN.CENTER)

# Save
out = r"c:\Rishabh\eduPulse\EduPulse_Presentation.pptx"
prs.save(out)
print(f"Presentation saved to {out}")
