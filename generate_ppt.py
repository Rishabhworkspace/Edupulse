from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Erudex"
    subtitle.text = "A Premium Learning Management System (MERN Stack)\nArchitecture, Design & Implementation"
    
    # 2. Project Overview
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Project Overview"
    tf = body_shape.text_frame
    tf.text = "Erudex is a comprehensive, scalable LMS platform."
    
    p = tf.add_paragraph()
    p.text = "Target Audience: Students, Instructors, and Administrators."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Core Features:"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "User Management System (RBAC)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Course Management System (Video & Curriculum)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Payment Management System (Stripe)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Coupon Management System"
    p.level = 2

    # 3. Technology Stack
    slide = prs.slides.add_slide(bullet_slide_layout)
    title_shape = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title_shape.text = "Technology Stack (MERN)"
    tf = body_shape.text_frame
    
    p = tf.add_paragraph()
    p.text = "Frontend"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "React.js 18 + Vite (Fast compilation)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Tailwind CSS (Custom UI tokens, Dark/Light Mode)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Redux Toolkit (State Management)"
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Backend & Database"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Node.js & Express.js (REST API)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "MongoDB & Mongoose (NoSQL Schemas)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Cloudinary (Media Hosting) & Stripe (Payments)"
    p.level = 1

    # 4. System Architecture
    slide = prs.slides.add_slide(bullet_slide_layout)
    title_shape = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title_shape.text = "System Architecture"
    tf = body_shape.text_frame
    
    p = tf.add_paragraph()
    p.text = "MVC Pattern on Backend"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Models: Defining data schema"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Controllers: Business logic"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Routes: API endpoints definition"
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Frontend Architecture"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Component-based UI with Lazy Loading"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Route Guards (RoleRoute, ProtectedRoute)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Axios Interceptors for JWT token rotation"
    p.level = 1

    # 5. Database Schema Design
    slide = prs.slides.add_slide(bullet_slide_layout)
    title_shape = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title_shape.text = "Database Schema Design"
    tf = body_shape.text_frame
    
    p = tf.add_paragraph()
    p.text = "User Schema: Handles multi-roles (student, instructor, admin), passwords, JWT."
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Course Schema: Master course details, nested curriculum structure."
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Enrollment Schema: Tracks student progress, videos watched, completion."
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Order Schema: Records financial transactions, Stripe intents, coupon applications."
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Coupon Schema: Usage limits, discount logic (flat vs percentage)."
    p.level = 0

    # 6. Security & Performance
    slide = prs.slides.add_slide(bullet_slide_layout)
    title_shape = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title_shape.text = "Security & Performance"
    tf = body_shape.text_frame
    
    p = tf.add_paragraph()
    p.text = "Authentication & Authorization"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Secure HttpOnly cookies / Access Tokens"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "bcryptjs for password hashing"
    p.level = 1

    p = tf.add_paragraph()
    p.text = "API Security"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Express Rate Limit & Helmet for header protection"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Data Sanitization against NoSQL injection & XSS"
    p.level = 1

    # 7. Conclusion
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Thank You"
    subtitle.text = "Erudex LMS is ready for deployment and scaling."

    prs.save('Erudex_Presentation.pptx')
    print("Erudex_Presentation.pptx generated successfully!")

if __name__ == '__main__':
    create_presentation()
